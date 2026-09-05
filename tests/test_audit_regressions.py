"""Regression tests for defects found during the full MCP audit."""

import json
import sqlite3
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from pypdf import PdfWriter

from samwoo_cm_bridge import server
from samwoo_cm_bridge.core.doc_cache_manager import DocumentCacheManager
from samwoo_cm_bridge.core.doc_parser import DocumentParser, SecurityError
from samwoo_cm_bridge.core.docx_exporter import DocxExporter
from samwoo_cm_bridge.core.equipment_quantity_auditor import EquipmentQuantityAuditor
from samwoo_cm_bridge.core.flexible_schedule_analyzer import FlexibleScheduleAnalyzer
from samwoo_cm_bridge.core.fire_hazard_conflict_detector import FireHazardConflictDetector
from samwoo_cm_bridge.core.formula_engine import safe_eval
from samwoo_cm_bridge.core.inspection_ncr_generator import InspectionNCRGenerator
from samwoo_cm_bridge.core.ocr_parser import MaterialCertOCRParser
from samwoo_cm_bridge.core.project_memory_engine import ProjectMemoryEngine
from samwoo_cm_bridge.core.video_record_manager import VideoRecordManager
from samwoo_cm_bridge.core.weather_stop_work_trigger import WeatherStopWorkTrigger


def test_cached_document_cannot_bypass_basename_validation(tmp_path):
    (tmp_path / "note.txt").write_text("검증할 문서 내용입니다.", encoding="utf-8")
    parser = DocumentParser(tmp_path)
    assert parser.parse_document("note.txt")["status"] == "SUCCESS"

    with pytest.raises(SecurityError):
        parser.parse_document("folder/../note.txt")


def test_table_only_docx_parses_without_unbound_line_number(tmp_path):
    path = tmp_path / "table-only.docx"
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "표 전용 문서"
    document.save(path)

    result = DocumentParser(tmp_path).parse_document(path.name)
    assert result["status"] == "SUCCESS"
    assert result["chunks"][0]["line_no"] == 1


def test_exporter_does_not_overwrite_existing_draft(tmp_path):
    exporter = DocxExporter(tmp_path)
    first = exporter.export("검토서.docx", "첫 번째 내용")
    second = exporter.export("검토서.docx", "두 번째 내용")

    assert first["docx_path"] != second["docx_path"]
    assert "첫 번째 내용" in Path(first["md_path"]).read_text(encoding="utf-8")
    assert "두 번째 내용" in Path(second["md_path"]).read_text(encoding="utf-8")


def test_cache_summary_path_cannot_escape_summary_directory(tmp_path):
    outside = tmp_path / "outside.txt"
    outside.write_text("secret", encoding="utf-8")
    manager = DocumentCacheManager(tmp_path / "data")

    assert manager._resolve_summary_path(str(outside)) is None


def test_formula_sandbox_rejects_resource_abuse_and_non_finite_values():
    with pytest.raises(ValueError, match="exponent"):
        safe_eval("2 ** 100", {})
    with pytest.raises(ValueError, match="Keyword"):
        safe_eval("round(1.234, ndigits=2)", {})
    with pytest.raises(ValueError, match="finite"):
        safe_eval("x + 1", {"x": float("inf")})
    with pytest.raises(ValueError, match="result"):
        safe_eval("1000000 * 1000000 * 1000000", {})


def test_server_formula_wrapper_rejects_bad_json_and_preserves_zero():
    malformed = json.loads(
        server.verify_calculation_safety(item_name="버팀보", variables_json="[]")
    )
    assert malformed["status"] == "ERROR"

    zero = json.loads(
        server.verify_calculation_safety(
            item_name="버팀보",
            design_val=10.0,
            allowable_val=0.0,
            req_sf=1.25,
        )
    )
    assert zero["status"] == "SUCCESS"
    assert zero["calculated_value"] == 0.0
    assert "FAIL" in zero["judgement"]

    zero_denominator = json.loads(
        server.verify_calculation_safety(
            item_name="버팀보",
            design_val=0.0,
            allowable_val=20.0,
            req_sf=1.25,
        )
    )
    assert zero_denominator["status"] == "ERROR"

    boolean_formula = json.loads(
        server.verify_calculation_safety(
            item_name="비교식",
            custom_formula="x > 1",
            variables_json='{"x": 2}',
        )
    )
    assert boolean_formula["status"] == "ERROR"


def test_text_replace_with_unequal_lengths_keeps_all_changes(tmp_path):
    from samwoo_cm_bridge.core.diff_audit_engine import DiffAuditEngine

    (tmp_path / "base.txt").write_text("첫 문장\n삭제될 안전 기준\n", encoding="utf-8")
    (tmp_path / "target.txt").write_text("바뀐 문장\n", encoding="utf-8")
    result = DiffAuditEngine(DocumentParser(tmp_path)).audit_text_diff(
        "base.txt", "target.txt"
    )

    assert result["total_changes"] == 2
    assert {change["change_type"] for change in result["changes"]} == {
        "MODIFIED",
        "DELETED",
    }


def test_project_memory_reindex_keeps_parent_and_replaces_actions(tmp_path):
    directive = tmp_path / "directive.txt"
    directive.write_text(
        "문서번호: TEST-001\n제목: 버팀보 변경 지시\n단면 증대 조치 요청\n",
        encoding="utf-8",
    )
    engine = ProjectMemoryEngine(parser=DocumentParser(tmp_path))

    engine.index_document(directive.name)
    engine.index_document(directive.name)

    with sqlite3.connect(engine.db_path) as conn:
        instructions = conn.execute("SELECT id FROM project_instructions").fetchall()
        actions = conn.execute("SELECT instruction_id FROM action_items").fetchall()
    assert len(instructions) == 1
    assert len(actions) == 1
    assert actions[0][0] == instructions[0][0]
    assert engine.search_memory("   ")["status"] == "ERROR"


def test_schedule_decimal_percentages_are_scaled(tmp_path):
    workbook_path = tmp_path / "schedule.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["공종", "계획공정률", "실적공정률"])
    ws.append(["골조공사", 0.85, 0.80])
    wb.save(workbook_path)

    result = FlexibleScheduleAnalyzer(DocumentParser(tmp_path)).analyze_schedule(
        workbook_path.name
    )
    item = result["all_activities"][0]
    assert item["planned_pct"] == 85.0
    assert item["actual_pct"] == 80.0
    assert item["variance_pct"] == -5.0


def test_schedule_percent_sign_is_not_scaled_twice(tmp_path):
    workbook_path = tmp_path / "schedule-percent.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["공종", "계획공정률", "실적공정률"])
    ws.append(["착공", "1%", "0%"])
    wb.save(workbook_path)

    result = FlexibleScheduleAnalyzer(DocumentParser(tmp_path)).analyze_schedule(
        workbook_path.name
    )
    item = result["all_activities"][0]
    assert item["planned_pct"] == 1.0
    assert item["actual_pct"] == 0.0


def test_ocr_only_passes_with_explicit_evidence(tmp_path):
    certificate = tmp_path / "certificate.txt"
    certificate.write_text(
        "성적서번호: TEST-001\n시료명: SS275\n항복강도: 300 MPa\n판정: 적합 (PASS)\n",
        encoding="utf-8",
    )

    result = MaterialCertOCRParser(DocumentParser(tmp_path)).parse_scanned_certificate(
        certificate.name
    )
    assert result["status"] == "SUCCESS"
    assert result["test_results"]["항복강도_MPa"] == 300.0
    assert result["ks_compliance_verdict"] == "적합 (PASS)"


def test_ocr_does_not_apply_unprovided_material_thresholds(tmp_path):
    certificate = tmp_path / "unjudged.txt"
    certificate.write_text(
        "성적서번호: TEST-002\n시료명: SS275\n항복강도: 200 MPa\n시험 결과 수치만 기재\n",
        encoding="utf-8",
    )

    result = MaterialCertOCRParser(DocumentParser(tmp_path)).parse_scanned_certificate(
        certificate.name
    )
    assert result["ks_compliance_verdict"] == "검증 필요 (REVIEW_REQUIRED)"


def test_empty_inspection_is_not_approved(tmp_path):
    generator = InspectionNCRGenerator(DocxExporter(tmp_path))
    empty = generator.generate_inspection_sheet("골조", "지하 1층")
    explicit = generator.generate_inspection_sheet(
        "골조",
        "지하 1층",
        inspection_items=[
            {
                "no": 1,
                "item": "철근 간격",
                "standard": "승인도서",
                "tolerance": "입력값",
                "contractor_check": "확인",
                "cm_verdict": "적합 (PASS)",
            }
        ],
    )
    assert "REVIEW_REQUIRED" in empty["final_verdict"]
    assert "PASS" in explicit["final_verdict"]
    assert "최종 승인 아님" in explicit["final_verdict"]


def test_blank_drawing_does_not_create_fake_equipment(tmp_path):
    pdf_path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with pdf_path.open("wb") as handle:
        writer.write(handle)

    auditor = EquipmentQuantityAuditor(parser=DocumentParser(tmp_path))
    with pytest.raises(ValueError, match="장비 일람표"):
        auditor.extract_pdf_equipment_tables(pdf_path.name)


def test_video_log_uses_parser_sandbox_and_requires_source_files(tmp_path):
    manager = VideoRecordManager(parser=DocumentParser(tmp_path))
    result = manager.generate_log(
        [{"video_file": "missing.mp4", "work_type": "철근 배근"}]
    )

    assert result["evidence_status"] == "REVIEW_REQUIRED"
    assert result["missing_video_files"] == 1
    assert Path(result["docx_path"]).parent == tmp_path.resolve()


def test_video_file_presence_does_not_verify_its_contents(tmp_path):
    (tmp_path / "evidence.mp4").write_bytes(b"not inspected by this tool")
    manager = VideoRecordManager(parser=DocumentParser(tmp_path))
    result = manager.generate_log(
        [{"video_file": "evidence.mp4", "work_type": "철근 배근"}]
    )

    assert result["records"][0]["file_status"] == "FOUND"
    assert result["source_file_status"] == "ALL_FOUND"
    assert result["evidence_status"] == "REVIEW_REQUIRED"


def test_fire_keywords_without_location_require_review(tmp_path):
    detector = FireHazardConflictDetector(parser=DocumentParser(tmp_path))
    result = detector.detect_conflicts(["용접 작업", "우레탄폼 시공"])

    assert result["status"] == "REVIEW_REQUIRED"
    assert result["conflicts"][0]["confidence"] == "POTENTIAL_UNKNOWN_LOCATION"


def test_weather_screening_does_not_invent_project_thresholds(tmp_path):
    trigger = WeatherStopWorkTrigger(parser=DocumentParser(tmp_path))
    missing_criteria = trigger.evaluate_and_issue_order(
        rain_mm=8.5,
        wind_speed_ms=13.0,
        planned_work="콘크리트 타설 및 타워크레인 양중",
    )
    below_supplied_criteria = trigger.evaluate_and_issue_order(
        rain_mm=8.5,
        wind_speed_ms=13.0,
        rain_threshold_mm=10.0,
        wind_threshold_ms=15.0,
        planned_work="콘크리트 타설 및 타워크레인 양중",
    )

    assert missing_criteria["status"] == "STOP_REVIEW_REQUIRED"
    assert all("미입력" in item["screening_criterion"] for item in missing_criteria["stop_items"])
    assert all("REVIEW_REQUIRED" in item["stop_level"] for item in missing_criteria["stop_items"])
    assert below_supplied_criteria["stop_items_count"] == 0


def test_previously_broken_server_wrappers_are_callable():
    indexed = json.loads(
        server.index_project_instruction("sample_발주처_설계변경지시공문.hwpx")
    )
    searched = json.loads(server.search_project_memory("버팀보"))
    tracked = json.loads(
        server.track_design_changes(
            "sample_설계변경_총괄내역서.xlsx",
            "sample_건축_단열및시공계획서.docx",
        )
    )

    assert indexed["status"] == "SUCCESS"
    assert searched["status"] == "SUCCESS"
    assert tracked["status"] == "SUCCESS"
