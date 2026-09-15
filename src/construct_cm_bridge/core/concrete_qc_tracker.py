"""Concrete Pouring & 28-day Compressive Strength QC Tracker (Module 19)

Tracks concrete pours, calculates 7-day / 28-day strength test due dates,
evaluates compressive strength compliance vs design standard (fck),
and accumulates test records into a persistent Excel Quality Control Ledger (XLSX).
"""

import os
import math
import logging
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, Any, List, Optional, Union

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .doc_parser import DocumentParser, SECURE_DATA_DIR
from .docx_exporter import DocxExporter

logger = logging.getLogger(__name__)

LEDGER_FILENAME = "콘크리트_품질관리대장.xlsx"


class ConcreteQCTracker:
    """Tracks concrete pouring logs and manages compressive strength quality control records."""

    def __init__(self, secure_dir: Optional[Union[str, Path]] = None):
        self.secure_dir = Path(secure_dir).resolve() if secure_dir else SECURE_DATA_DIR.resolve()
        self.secure_dir.mkdir(parents=True, exist_ok=True)
        self.ledger_path = self.secure_dir / LEDGER_FILENAME
        self._ensure_ledger_exists()

    def _ensure_ledger_exists(self):
        """Initializes the Excel QC ledger if not already present."""
        if not self.ledger_path.exists():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "콘크리트품질관리대장"

            # Header row styling
            headers = [
                "타설번호", "타설일자", "타설부위/구조체", "설계기준강도(fck, MPa)", "타설량(m3)",
                "레미콘 규격", "7일 시험일자", "7일 강도(MPa)", "7일 발현율(%)",
                "28일 시험일자", "28일 강도(MPa)", "28일 발현율(%)", "최종 판정", "비고/감리확인"
            ]
            ws.append(headers)

            navy_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
            white_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
            center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for col_idx, cell in enumerate(ws[1], 1):
                cell.fill = navy_fill
                cell.font = white_font
                cell.alignment = center_align

            ws.column_dimensions["A"].width = 12
            ws.column_dimensions["B"].width = 14
            ws.column_dimensions["C"].width = 22
            ws.column_dimensions["D"].width = 18
            ws.column_dimensions["E"].width = 14
            ws.column_dimensions["F"].width = 18
            ws.column_dimensions["G"].width = 14
            ws.column_dimensions["H"].width = 14
            ws.column_dimensions["I"].width = 14
            ws.column_dimensions["J"].width = 14
            ws.column_dimensions["K"].width = 14
            ws.column_dimensions["L"].width = 14
            ws.column_dimensions["M"].width = 14
            ws.column_dimensions["N"].width = 20

            wb.save(self.ledger_path)

    def register_pour(
        self,
        date_str: str,
        location: str,
        spec_fck: float,
        volume_m3: float,
        remicon_spec: str = "미입력",
        measured_7d_mpa: Optional[float] = None,
        measured_28d_mpa: Optional[float] = None,
        project_name: str = "미입력 프로젝트",
    ) -> Dict[str, Any]:
        """Registers a concrete pour event, computes 7/28d dates, checks strength, and updates ledger."""
        try:
            pour_date = datetime.strptime(date_str.replace(".", "-").replace("/", "-").strip(), "%Y-%m-%d")
        except ValueError as e:
            raise ValueError("date_str은 유효한 YYYY-MM-DD, YYYY.MM.DD 또는 YYYY/MM/DD 형식이어야 합니다.") from e

        numeric_values = {
            "spec_fck": spec_fck,
            "volume_m3": volume_m3,
            "measured_7d_mpa": measured_7d_mpa,
            "measured_28d_mpa": measured_28d_mpa,
        }
        for name, value in numeric_values.items():
            if value is None:
                continue
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(float(value)):
                raise ValueError(f"{name}은 유한한 숫자여야 합니다.")
            if value < 0:
                raise ValueError(f"{name}은 음수일 수 없습니다.")
        if spec_fck <= 0:
            raise ValueError("spec_fck는 0보다 커야 합니다.")

        d7_date = pour_date + timedelta(days=7)
        d28_date = pour_date + timedelta(days=28)

        d7_str = d7_date.strftime("%Y.%m.%d")
        d28_str = d28_date.strftime("%Y.%m.%d")

        # Strength calculation & compliance check
        rate_7d = (measured_7d_mpa / spec_fck * 100.0) if measured_7d_mpa is not None else None
        rate_28d = (measured_28d_mpa / spec_fck * 100.0) if measured_28d_mpa is not None else None

        verdict = "시험 대기 (PENDING)"
        if measured_28d_mpa is not None:
            if measured_28d_mpa >= spec_fck:
                verdict = "28일 강도 적합 (PASS)"
            else:
                verdict = "28일 강도 미달 (FAIL)"
        elif measured_7d_mpa is not None:
            if measured_7d_mpa >= spec_fck * 0.65:
                verdict = "7일 강도 양호 (NORMAL)"
            else:
                verdict = "7일 강도 주의 (WARNING)"

        # Append to Excel Ledger
        wb = openpyxl.load_workbook(self.ledger_path)
        ws = wb["콘크리트품질관리대장"]
        row_count = ws.max_row
        pour_no = f"CONC-{pour_date.strftime('%Y%m%d')}-{row_count:02d}"

        row_data = [
            pour_no,
            pour_date.strftime("%Y.%m.%d"),
            location,
            spec_fck,
            volume_m3,
            remicon_spec,
            d7_str,
            measured_7d_mpa if measured_7d_mpa is not None else "-",
            f"{rate_7d:.1f}%" if rate_7d is not None else "-",
            d28_str,
            measured_28d_mpa if measured_28d_mpa is not None else "-",
            f"{rate_28d:.1f}%" if rate_28d is not None else "-",
            verdict,
            "입력 자료 기반 등록",
        ]
        ws.append(row_data)

        # Style new row
        r_idx = ws.max_row
        for c in ws[r_idx]:
            c.font = Font(name="맑은 고딕", size=9.5)
            c.alignment = Alignment(horizontal="center", vertical="center")

        wb.save(self.ledger_path)

        # Compute days until tests
        today = datetime.now().date()
        days_to_7d = (d7_date.date() - today).days
        days_to_28d = (d28_date.date() - today).days

        return {
            "status": "SUCCESS",
            "pour_no": pour_no,
            "pour_date": pour_date.strftime("%Y.%m.%d"),
            "location": location,
            "spec_fck_mpa": spec_fck,
            "volume_m3": volume_m3,
            "remicon_spec": remicon_spec,
            "test_7d_date": d7_str,
            "days_until_7d_test": days_to_7d,
            "measured_7d_mpa": measured_7d_mpa,
            "rate_7d_pct": rate_7d,
            "test_28d_date": d28_str,
            "days_until_28d_test": days_to_28d,
            "measured_28d_mpa": measured_28d_mpa,
            "rate_28d_pct": rate_28d,
            "verdict": verdict,
            "ledger_path": str(self.ledger_path),
            "source_anchor": f"{LEDGER_FILENAME} [콘크리트품질관리대장!R{r_idx}]",
        }


# Singleton instance
_concrete_tracker = ConcreteQCTracker()


def register_concrete_pour(
    date_str: str,
    location: str,
    spec_fck: float,
    volume_m3: float,
    remicon_spec: str = "미입력",
    measured_7d_mpa: Optional[float] = None,
    measured_28d_mpa: Optional[float] = None,
    project_name: str = "미입력 프로젝트",
) -> Dict[str, Any]:
    return _concrete_tracker.register_pour(
        date_str=date_str,
        location=location,
        spec_fck=spec_fck,
        volume_m3=volume_m3,
        remicon_spec=remicon_spec,
        measured_7d_mpa=measured_7d_mpa,
        measured_28d_mpa=measured_28d_mpa,
        project_name=project_name,
    )
