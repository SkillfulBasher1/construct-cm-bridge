"""Samwoo CM Standard Review Document Exporter (Module 3 - Exporter)

Generates professional Samwoo CM style inspection/review reports in both Word (.docx) and Markdown (.md).
Implements standard 4-column 3-way cross examination tables, formula verification breakdowns,
and official CM signature blocks.
"""

import os
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Optional, List, Union

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from .doc_parser import SECURE_DATA_DIR


def set_cell_background(cell, fill_hex: str):
    """Sets background color of a Word table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner padding for table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


class DocxExporter:
    """Exports structured CM Review Opinions into Docx and Markdown."""

    def __init__(self, output_dir: Optional[Union[str, Path]] = None):
        self.output_dir = Path(output_dir).resolve() if output_dir else SECURE_DATA_DIR.resolve()
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(
        self,
        output_filename: str,
        report_text: str,
        project_name: Optional[str] = "삼우씨엠 신축공사 CM현장",
        reviewer_name: Optional[str] = "수석 건설사업관리기술인",
        discipline: Optional[str] = "토목 / 구조 / 기계 / 소방 / 전기",
        doc_no: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Generates both .docx and .md review documents."""
        base_name = os.path.splitext(os.path.basename(output_filename))[0]
        docx_path = self.output_dir / f"{base_name}.docx"
        md_path = self.output_dir / f"{base_name}.md"

        now_str = datetime.now().strftime("%Y년 %m월 %d일")
        doc_number = doc_no or f"SWCM-REV-{datetime.now().strftime('%Y%m%d')}-01"

        # 1. Write Markdown file
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(f"# [건설사업관리(CM) 기술검토의견서]\n\n")
            f.write(f"- **문서번호:** {doc_number}\n")
            f.write(f"- **프로젝트명:** {project_name}\n")
            f.write(f"- **검토분야:** {discipline}\n")
            f.write(f"- **검토일자:** {now_str}\n")
            f.write(f"- **검토자:** {reviewer_name}\n\n")
            f.write(f"---\n\n")
            f.write(report_text)
            f.write(f"\n\n---\n**주식회사 삼우씨엠건축사사무소 건설사업관리단**\n")

        # 2. Build Word (.docx) document with styling
        doc = docx.Document()

        # Set standard margins (20mm ~ 0.79 in)
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run("건설사업관리(CM) 기술검토의견서")
        title_run.font.name = "맑은 고딕"
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)  # Samwoo Deep Navy

        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run("Technical Review & Cross-Examination Report")
        sub_run.font.name = "Arial"
        sub_run.font.size = Pt(10)
        sub_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

        # Meta Header Table
        meta_table = doc.add_table(rows=3, cols=4)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False

        meta_data = [
            [("문서번호", doc_number), ("검토일자", now_str)],
            [("프로젝트명", project_name), ("검토분야", discipline)],
            [("검토자", reviewer_name), ("관리기관", "㈜삼우씨엠건축사사무소")],
        ]

        for r_idx, row_pairs in enumerate(meta_data):
            # Col 0, 1
            lbl1, val1 = row_pairs[0]
            lbl2, val2 = row_pairs[1]

            cell_lbl1 = meta_table.cell(r_idx, 0)
            cell_val1 = meta_table.cell(r_idx, 1)
            cell_lbl2 = meta_table.cell(r_idx, 2)
            cell_val2 = meta_table.cell(r_idx, 3)

            cell_lbl1.text = lbl1
            cell_val1.text = val1
            cell_lbl2.text = lbl2
            cell_val2.text = val2

            for c in [cell_lbl1, cell_lbl2]:
                set_cell_background(c, "EDF2F7")
                set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

            for c in [cell_val1, cell_val2]:
                set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                p = c.paragraphs[0]
                for run in p.runs:
                    run.font.size = Pt(9.5)

        doc.add_paragraph()  # Spacer

        # Parse and render Markdown body lines into docx
        self._render_markdown_body_to_docx(doc, report_text)

        # Footer Signature Box
        doc.add_paragraph()
        sig_p = doc.add_paragraph()
        sig_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        sig_run = sig_p.add_run(f"작성일: {now_str}\n주식회사 삼우씨엠건축사사무소 건설사업관리단 (인)")
        sig_run.font.name = "맑은 고딕"
        sig_run.font.size = Pt(10)
        sig_run.font.bold = True
        sig_run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

        # Save docx
        doc.save(str(docx_path))

        return {
            "status": "SUCCESS",
            "docx_path": str(docx_path),
            "md_path": str(md_path),
            "filename": docx_path.name,
            "project_name": project_name,
            "doc_number": doc_number,
        }

    def _render_markdown_body_to_docx(self, doc: docx.Document, text: str):
        """Converts Markdown headings, paragraphs, and tables into styled Word elements."""
        lines = text.strip().split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].strip()

            if not line:
                i += 1
                continue

            # Headings
            if line.startswith("# "):
                h = doc.add_heading(level=1)
                run = h.add_run(line[2:].strip())
                run.font.name = "맑은 고딕"
                run.font.size = Pt(14)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
                i += 1
            elif line.startswith("## "):
                h = doc.add_heading(level=2)
                run = h.add_run(line[3:].strip())
                run.font.name = "맑은 고딕"
                run.font.size = Pt(12)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
                i += 1
            elif line.startswith("### "):
                h = doc.add_heading(level=3)
                run = h.add_run(line[4:].strip())
                run.font.name = "맑은 고딕"
                run.font.size = Pt(10.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
                i += 1
            # Markdown Table detection
            elif line.startswith("|") and line.endswith("|"):
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                self._render_markdown_table_to_docx(doc, table_lines)
            # Bullet list
            elif line.startswith("- ") or line.startswith("* "):
                p = doc.add_paragraph(style="List Bullet")
                self._add_formatted_text(p, line[2:].strip())
                i += 1
            # Numbered list
            elif re.match(r"^\d+\.\s", line):
                match = re.match(r"^\d+\.\s", line)
                p = doc.add_paragraph(style="List Number")
                self._add_formatted_text(p, line[match.end():].strip())
                i += 1
            # Regular paragraph
            else:
                p = doc.add_paragraph()
                self._add_formatted_text(p, line)
                i += 1

    def _add_formatted_text(self, paragraph, text: str):
        """Helper to parse **bold** and inline text formatting."""
        parts = re.split(r"(\*\*.*?\*\*)", text)
        for part in parts:
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            else:
                run = paragraph.add_run(part)
            run.font.name = "맑은 고딕"
            run.font.size = Pt(9.5)

    def _render_markdown_table_to_docx(self, doc: docx.Document, table_lines: List[str]):
        """Renders parsed markdown table lines with professional Samwoo CM styling."""
        if not table_lines:
            return

        rows_data = []
        for line in table_lines:
            # Skip delimiter line like |---|---|
            if set(line.replace("|", "").strip()) <= {"-", ":"}:
                continue
            cells = [c.strip() for c in line.strip("|").split("|")]
            rows_data.append(cells)

        if not rows_data:
            return

        cols_count = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=cols_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row_cells in enumerate(rows_data):
            for c_idx in range(cols_count):
                cell = table.cell(r_idx, c_idx)
                cell_text = row_cells[c_idx] if c_idx < len(row_cells) else ""
                cell.text = cell_text

                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)

                if r_idx == 0:
                    # Header row styling (Deep Navy)
                    set_cell_background(cell, "1A365D")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.bold = True
                        run.font.name = "맑은 고딕"
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    # Alternating row background or status badges
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F7FAFC")

                    # Highlight PASS / FAIL badges and Source Anchors
                    for run in p.runs:
                        run.font.name = "맑은 고딕"
                        run.font.size = Pt(9.0)
                        if "PASS" in run.text or "적합" in run.text or "RECEIVED" in run.text:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0x27, 0x67, 0x49)  # Green
                        elif "FAIL" in run.text or "부적합" in run.text or "MISSING" in run.text:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xC5, 0x30, 0x30)  # Red
                        elif "출처" in run.text or "L" in run.text and ("char" in run.text or "Table" in run.text or "R" in run.text):
                            run.font.size = Pt(8.0)
                            run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)  # Slate Gray

        doc.add_paragraph()  # Spacer


# Singleton instance
_exporter = DocxExporter()


def export_review_document(
    output_filename: str,
    report_text: str,
    project_name: Optional[str] = "삼우씨엠 신축공사 CM현장",
    reviewer_name: Optional[str] = "수석 건설사업관리기술인",
    discipline: Optional[str] = "토목 / 구조 / 기계 / 소방 / 전기",
    doc_no: Optional[str] = None,
) -> Dict[str, Any]:
    return _exporter.export(
        output_filename=output_filename,
        report_text=report_text,
        project_name=project_name,
        reviewer_name=reviewer_name,
        discipline=discipline,
        doc_no=doc_no,
    )
