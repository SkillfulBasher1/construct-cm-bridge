"""CM Inspection Sheet & Non-Conformance Report (NCR) Generator (Module 15 - Inspection & NCR)

Generates:
1. Standard CM Inspection Request & Result Sheet (검측요청서 및 결과통보서)
2. Non-Conformance Report & Corrective Action Order (부적합 시정지시서 - NCR)
- User-entered defects, evidence descriptions, and a review-required action draft.
"""

import os
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional, Union

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from .docx_exporter import DocxExporter, choose_output_stem, set_cell_background, set_cell_margins


class InspectionNCRGenerator:
    """Generates Inspection Request/Result Sheets and Non-Conformance Orders (NCR)."""

    def __init__(self, exporter: Optional[DocxExporter] = None):
        self.exporter = exporter or DocxExporter()
        self.output_dir = self.exporter.output_dir

    def generate_inspection_sheet(
        self,
        work_type: str,
        location: str,
        contractor_spec: str = "",
        inspection_items: Optional[List[Dict[str, str]]] = None,
        project_name: str = "미입력 프로젝트",
        inspector_name: str = "미입력 분야별 책임기술인",
    ) -> Dict[str, Any]:
        """Generates standard CM Inspection Request & Result Sheet."""
        now = datetime.now()
        date_str = now.strftime("%Y.%m.%d")
        doc_no = f"SWCM-INSP-{now.strftime('%Y%m%d')}-01"

        items = list(inspection_items or [])

        has_fail = any("부적합" in i.get("cm_verdict", "") or "FAIL" in i.get("cm_verdict", "") for i in items)
        all_pass = bool(items) and all(
            "적합" in i.get("cm_verdict", "") or "PASS" in i.get("cm_verdict", "")
            for i in items
        )
        if has_fail:
            final_verdict = "검측 부적합 (FAIL)"
            cm_opinion = "입력된 부적합 항목을 보완하고 재검측해야 합니다."
        elif all_pass:
            final_verdict = "입력 판정 전 항목 PASS (최종 승인 아님)"
            cm_opinion = "입력된 모든 검측 항목이 PASS입니다. 원본 증빙과 승인권자 서명은 별도 확인해야 합니다."
        else:
            final_verdict = "검측 판정 대기 (REVIEW_REQUIRED)"
            cm_opinion = "검측 항목과 감리원 판정이 입력되지 않아 승인할 수 없습니다."

        # Build Markdown content
        md_lines = [
            f"# [검측요청서 및 결과통보서]",
            f"- **문서번호:** {doc_no}",
            f"- **공 사 명:** {project_name}",
            f"- **검측일자:** {date_str}",
            f"- **검측공종:** {work_type}",
            f"- **검측위치:** {location}",
            f"- **시공사 사양:** {contractor_spec or '미입력'}",
            f"- **담당 감리원:** {inspector_name}\n",
            f"# 1. 세부 검측 체크리스트 및 판정 결과",
            f"| No | 검측 세부 항목 | 관리 기준 | 허용 오차 | 시공사 점검 | **감리원 판정** |",
            f"|---|---|---|---|---|---|",
        ]

        for it in items:
            md_lines.append(
                f"| {it.get('no')} | {it.get('item')} | {it.get('standard')} | {it.get('tolerance')} | {it.get('contractor_check')} | **{it.get('cm_verdict')}** |"
            )

        md_lines.extend([
            f"\n# 2. 종합 검측 판정 및 지시사항",
            f"- **최종 검측 결과:** **{final_verdict}**",
            f"- **감리의견:** {cm_opinion}",
        ])

        report_md = "\n".join(md_lines)
        output_filename = f"검측결과통보서_{work_type}_{now.strftime('%Y%m%d')}.docx"

        res = self.exporter.export(
            output_filename=output_filename,
            report_text=report_md,
            project_name=project_name,
            reviewer_name=inspector_name,
            discipline=f"검측업무 ({work_type})",
            doc_no=doc_no,
        )

        return {
            "status": "SUCCESS",
            "doc_no": doc_no,
            "work_type": work_type,
            "location": location,
            "final_verdict": final_verdict,
            "docx_path": res.get("docx_path"),
            "md_path": res.get("md_path"),
        }

    def draft_ncr_order(
        self,
        issue_description: str,
        location: str,
        defect_category: str = "미분류",
        photo_attached: str = "미첨부",
        corrective_deadline: str = "",
        recipient: str = "미입력 시공사 현장소장",
        project_name: str = "미입력 프로젝트",
        chief_cm_name: str = "미입력 책임기술인",
    ) -> Dict[str, Any]:
        """Drafts formal Non-Conformance Report (부적합 시정지시서 - NCR)."""
        now = datetime.now()
        date_str = now.strftime("%Y년 %m월 %d일")
        deadline_str = corrective_deadline or "미입력"
        doc_no = f"SWCM-NCR-{now.strftime('%Y%m%d')}-01"

        md_content = f"""# 주식회사 삼우씨엠건축사사무소
**부적합 시정지시서 초안 (NCR DRAFT - 미승인)**

---
- **관리번호:** {doc_no}
- **발행일자:** {date_str}
- **수　　신:** {recipient}
- **발　　신:** {project_name} 책임건설사업관리기술인 {chief_cm_name}
- **부적합 유형:** **{defect_category}** (적용 기준 원문 별도 확인)
- **발생위치:** {location}

---

### 1. 부적합 및 지적 사항 내용
{issue_description}

### 2. 현장 실측/사진 증빙 현황
- **현장 증빙:** {photo_attached}
- **입력된 지적 내용:** {issue_description}
- **판정 상태:** 적용 기준·현장 상태·사진 원본 확인 필요 (REVIEW_REQUIRED)

### 3. 감리단 시정 및 조치 지시사항
1. 책임기술인은 입력된 지적 내용과 적용 설계도서·시방서·현장 상태를 확인할 것.
2. 부적합이 확인되면 후속공정 중지 범위, 시정조치 계획 및 재검측 절차를 확정할 것.
3. 확인 결과와 승인권자의 서명 후 정식 NCR을 발행할 것.

### 4. 시정조치 완료 보고 기한
- **이행 기한:** **{deadline_str}**

---
**초안 작성본 - 책임기술인 승인 서명 필요**
"""
        output_stem = choose_output_stem(self.output_dir, f"시정지시서_{doc_no}")
        output_docx = f"{output_stem}.docx"
        output_md = f"{output_stem}.md"
        docx_path = self.output_dir / output_docx
        md_path = self.output_dir / output_md

        with open(md_path, "w", encoding="utf-8") as f:
            f.write(md_content)

        doc = docx.Document()
        for s in doc.sections:
            s.top_margin = Inches(0.8)
            s.bottom_margin = Inches(0.8)
            s.left_margin = Inches(0.8)
            s.right_margin = Inches(0.8)

        # Header Title
        h_p = doc.add_paragraph()
        h_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h_run = h_p.add_run("부적합 시정지시서 초안 (NCR DRAFT)")
        h_run.font.name = "맑은 고딕"
        h_run.font.size = Pt(18)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(0x9B, 0x2C, 0x2C)  # Red accent for NCR

        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(f"주식회사 삼우씨엠건축사사무소 | {project_name}")
        sub_run.font.name = "맑은 고딕"
        sub_run.font.size = Pt(11)
        sub_run.font.bold = True
        sub_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

        # Meta table
        meta_table = doc.add_table(rows=4, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_data = [
            ("관리번호", doc_no),
            ("발행일자", date_str),
            ("부적합 유형", f"{defect_category} (발생위치: {location})"),
            ("시정기한", deadline_str),
        ]
        for r_idx, (lbl, val) in enumerate(meta_data):
            c0 = meta_table.cell(r_idx, 0)
            c1 = meta_table.cell(r_idx, 1)
            c0.text = lbl
            c1.text = val
            set_cell_background(c0, "FED7D7")  # Red tint for NCR
            set_cell_margins(c0, 60, 60, 80, 80)
            set_cell_margins(c1, 60, 60, 80, 80)
            c0.paragraphs[0].runs[0].font.bold = True

        doc.add_paragraph()

        # Body paragraphs
        p1 = doc.add_paragraph()
        p1_r = p1.add_run(f"1. 부적합 및 지적 사항\n   {issue_description}")
        p1_r.font.size = Pt(11)

        p2 = doc.add_paragraph()
        p2_r = p2.add_run(f"2. 현장 증빙 및 결함 상태\n   - 증빙: {photo_attached}\n   - 적용 기준: 승인 설계도서·시방서 원문 별도 확인")
        p2_r.font.size = Pt(11)

        p3 = doc.add_paragraph()
        p3_r = p3.add_run("3. 확인 및 승인 절차\n   가. 입력 지적사항과 원본 증빙 확인\n   나. 적용 설계도서·시방서와 대조\n   다. 책임기술인 승인 후 정식 NCR 발행")
        p3_r.font.size = Pt(11)

        doc.add_paragraph()
        sig_p = doc.add_paragraph()
        sig_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sig_run = sig_p.add_run(f"초안 검토자: {chief_cm_name}\n책임기술인 승인 서명 필요")
        sig_run.font.name = "맑은 고딕"
        sig_run.font.size = Pt(12)
        sig_run.font.bold = True
        sig_run.font.color.rgb = RGBColor(0x9B, 0x2C, 0x2C)

        doc.save(str(docx_path))

        return {
            "status": "SUCCESS",
            "doc_no": doc_no,
            "defect_category": defect_category,
            "location": location,
            "deadline": deadline_str,
            "recipient": recipient,
            "evidence_status": "REVIEW_REQUIRED",
            "docx_path": str(docx_path),
            "md_path": str(md_path),
        }


# Singleton instance
_inspection_ncr_gen = InspectionNCRGenerator()


def generate_inspection_sheet(
    work_type: str,
    location: str,
    contractor_spec: str = "",
    inspection_items: Optional[List[Dict[str, str]]] = None,
) -> Dict[str, Any]:
    return _inspection_ncr_gen.generate_inspection_sheet(
        work_type=work_type,
        location=location,
        contractor_spec=contractor_spec,
        inspection_items=inspection_items,
    )


def draft_ncr_correction_order(
    issue_description: str,
    location: str,
    defect_category: str = "미분류",
    photo_attached: str = "미첨부",
    corrective_deadline: str = "",
    recipient: str = "미입력 시공사 현장소장",
) -> Dict[str, Any]:
    return _inspection_ncr_gen.draft_ncr_order(
        issue_description=issue_description,
        location=location,
        defect_category=defect_category,
        photo_attached=photo_attached,
        corrective_deadline=corrective_deadline,
        recipient=recipient,
    )
