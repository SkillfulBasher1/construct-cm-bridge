"""NCR Corrective Action Before / After Photo Confirmation Sheet Builder (Module 20)

Generates official 'Before / After Corrective Action Verification Sheet (.docx / .md)'
mapping defect photos (Before) and contractor rectification photos (After) in a 1:1 side-by-side card layout.
"""

import os
import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional, Union

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

from .doc_parser import SECURE_DATA_DIR
from .docx_exporter import DocxExporter, set_cell_background, set_cell_margins

logger = logging.getLogger(__name__)


class NCRActionSheetBuilder:
    """Builds side-by-side Before/After photo comparison sheets."""

    def __init__(self, secure_dir: Optional[Union[str, Path]] = None):
        self.secure_dir = Path(secure_dir).resolve() if secure_dir else SECURE_DATA_DIR.resolve()
        self.secure_dir.mkdir(parents=True, exist_ok=True)
        self.exporter = DocxExporter(self.secure_dir)

    def generate_sheet(
        self,
        issue_title: str,
        before_img: str,
        after_img: str,
        description: str,
        location: str = "지하 2층 1구역",
        action_date: str = "",
        ncr_no: str = "SWCM-NCR-20260829-01",
        project_name: str = "삼우씨엠 신축공사 CM현장",
        inspector_name: str = "김수석 책임건설사업관리기술인",
        contractor_name: str = "(주)대우건설 현장소장",
    ) -> Dict[str, Any]:
        """Generates Before/After photo confirmation document."""
        now = datetime.now()
        date_str = action_date if action_date else now.strftime("%Y.%m.%d")
        doc_no = f"SWCM-ACT-{now.strftime('%Y%m%d')}-01"

        # 1. Generate Markdown
        md_lines = [
            f"# [현장 시정조치 완료 확인서 (Before / After)]",
            f"- **문서번호:** {doc_no} (관련 NCR: {ncr_no})",
            f"- **공 사 명:** {project_name}",
            f"- **지적 위치:** {location}",
            f"- **조치 일자:** {date_str}",
            f"- **감리 확인:** {inspector_name} / **시공사:** {contractor_name}\n",
            f"# 1. 지적사항 및 시정조치 개요",
            f"- **지적 건명:** **{issue_title}**",
            f"- **상세 내용:** {description}\n",
            f"# 2. 조치 전 / 후 (Before & After) 대조 사진대지",
            f"| 구분 | [지적 사항] 조치 전 (Before) | [조치 완료] 시정 후 (After) |",
            f"|---|---|---|",
            f"| 사진 파일 | `{before_img}` | `{after_img}` |",
            f"| 현장 상태 | 결함 발생 및 기준 미달 상태 | 시정조치 완료 및 감리원 재검측 합격 |",
            f"| 판정 | **시정지시 발부 (NCR)** | **시정완료 승인 (APPROVED)** |\n",
            f"# 3. 감리단 최종 검측 의견",
            f"- 상기 지적사항에 대하여 시공사가 제출한 조치내용 및 현장 실물을 확인한 결과, 설계도서 및 KCS 시공기준에 부합하게 적정 조치되었음을 확인함.",
        ]
        report_md = "\n".join(md_lines)

        # 2. Build Docx Document
        doc = docx.Document()
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Title
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run("현장 시정조치 완료 확인서 (Before / After)")
        title_run.font.name = "맑은 고딕"
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

        # Header Info Table
        info_table = doc.add_table(rows=3, cols=4)
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        info_data = [
            [("문서번호", doc_no), ("관련 NCR", ncr_no)],
            [("공사명", project_name), ("조치위치", location)],
            [("조치일자", date_str), ("검측확인", inspector_name)],
        ]
        for r_idx, row_pairs in enumerate(info_data):
            cell_lbl1, cell_val1 = info_table.cell(r_idx, 0), info_table.cell(r_idx, 1)
            cell_lbl2, cell_val2 = info_table.cell(r_idx, 2), info_table.cell(r_idx, 3)
            cell_lbl1.text, cell_val1.text = row_pairs[0]
            cell_lbl2.text, cell_val2.text = row_pairs[1]
            for c in [cell_lbl1, cell_lbl2]:
                set_cell_background(c, "EDF2F7")
                set_cell_margins(c, 70, 70, 90, 90)
                p = c.paragraphs[0]
                p.runs[0].font.bold = True
                p.runs[0].font.name = "맑은 고딕"
                p.runs[0].font.size = Pt(9.5)
            for c in [cell_val1, cell_val2]:
                set_cell_margins(c, 70, 70, 90, 90)
                p = c.paragraphs[0]
                p.runs[0].font.name = "맑은 고딕"
                p.runs[0].font.size = Pt(9.5)

        doc.add_paragraph()  # Spacer

        # 2-Column Photo Comparison Card
        p_card = doc.add_heading("2. 조치 전 / 후 (Before & After) 대조 사진", level=2)
        p_card.runs[0].font.name = "맑은 고딕"
        p_card.runs[0].font.size = Pt(12)
        p_card.runs[0].font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

        card_table = doc.add_table(rows=3, cols=2)
        card_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Headers
        cell_b_head = card_table.cell(0, 0)
        cell_a_head = card_table.cell(0, 1)
        cell_b_head.text = "【 조치 전 (BEFORE) 】"
        cell_a_head.text = "【 조치 후 (AFTER) 】"
        set_cell_background(cell_b_head, "9B2C2C")  # Dark Red
        set_cell_background(cell_a_head, "276749")  # Dark Green
        for c in [cell_b_head, cell_a_head]:
            set_cell_margins(c, 90, 90, 100, 100)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            p.runs[0].font.size = Pt(11)

        # Image Cells
        cell_b_img = card_table.cell(1, 0)
        cell_a_img = card_table.cell(1, 1)
        for c in [cell_b_img, cell_a_img]:
            set_cell_margins(c, 100, 100, 100, 100)

        before_path = (self.secure_dir / os.path.basename(before_img)).resolve()
        after_path = (self.secure_dir / os.path.basename(after_img)).resolve()

        p_b = cell_b_img.paragraphs[0]
        p_b.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if before_path.exists() and before_path.suffix.lower() in [".jpg", ".jpeg", ".png"]:
            try:
                p_b.add_run().add_picture(str(before_path), width=Inches(3.0))
            except Exception:
                p_b.add_run(f"📷 [지적 사진]\n{before_img}").font.size = Pt(9.5)
        else:
            p_b.add_run(f"📷 [지적 사진 (Before)]\n({before_img})").font.size = Pt(9.5)

        p_a = cell_a_img.paragraphs[0]
        p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if after_path.exists() and after_path.suffix.lower() in [".jpg", ".jpeg", ".png"]:
            try:
                p_a.add_run().add_picture(str(after_path), width=Inches(3.0))
            except Exception:
                p_a.add_run(f"📷 [조치 사진]\n{after_img}").font.size = Pt(9.5)
        else:
            p_a.add_run(f"📷 [조치 완료 사진 (After)]\n({after_img})").font.size = Pt(9.5)

        # Descriptions
        cell_b_desc = card_table.cell(2, 0)
        cell_a_desc = card_table.cell(2, 1)
        set_cell_background(cell_b_desc, "FFF5F5")
        set_cell_background(cell_a_desc, "F0FFF4")
        for c in [cell_b_desc, cell_a_desc]:
            set_cell_margins(c, 80, 80, 100, 100)

        cell_b_desc.text = f"• 지적사항: {issue_title}\n• 결함내용: {description[:80]}"
        cell_a_desc.text = f"• 조치결과: 현장 수정 및 재시공 완료\n• 감리확인: KCS 기준 적합 판정"

        doc.add_paragraph()  # Spacer

        # Bottom Approval Box
        sig_p = doc.add_paragraph()
        sig_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        sig_run = sig_p.add_run(f"확인일자: {date_str}\n책임건설사업관리기술인: {inspector_name} (서명/인)")
        sig_run.font.name = "맑은 고딕"
        sig_run.font.size = Pt(10)
        sig_run.font.bold = True
        sig_run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

        docx_filename = f"시정조치확인서_{now.strftime('%Y%m%d')}.docx"
        docx_path = self.secure_dir / docx_filename
        doc.save(str(docx_path))

        md_filename = f"시정조치확인서_{now.strftime('%Y%m%d')}.md"
        md_path = self.secure_dir / md_filename
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(report_md)

        return {
            "status": "SUCCESS",
            "doc_no": doc_no,
            "ncr_no": ncr_no,
            "issue_title": issue_title,
            "location": location,
            "action_date": date_str,
            "docx_path": str(docx_path),
            "md_path": str(md_path),
            "source_anchor": f"{docx_filename} [Before/After 카드]",
        }


# Singleton instance
_action_builder = NCRActionSheetBuilder()


def generate_before_after_sheet(
    issue_title: str,
    before_img: str,
    after_img: str,
    description: str,
    location: str = "지하 2층 1구역",
    action_date: str = "",
    ncr_no: str = "SWCM-NCR-20260829-01",
    project_name: str = "삼우씨엠 신축공사 CM현장",
    inspector_name: str = "김수석 책임건설사업관리기술인",
    contractor_name: str = "(주)대우건설 현장소장",
) -> Dict[str, Any]:
    return _action_builder.generate_sheet(
        issue_title=issue_title,
        before_img=before_img,
        after_img=after_img,
        description=description,
        location=location,
        action_date=action_date,
        ncr_no=ncr_no,
        project_name=project_name,
        inspector_name=inspector_name,
        contractor_name=contractor_name,
    )
